"""Read a SAVI v1.10 workbook (.xlsm) into tidy pandas tables.

Two things come out of a SAVI workbook:

* the **rule set** - SAVI's reference sheets (Ref00 ... Ref15), i.e. the SAMPT
  Part C tables indexed to a price year; and
* the **inventory** - the user's Structures / Elements / Planned / Routine
  input sheets plus the Main Page settings.

Only cached cell *values* are read (``data_only=True``); no Excel formulas are
evaluated here. Cell coordinates follow the named ranges defined in SAVI v1.10.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string as col

# Order of the six traffic-management arrangements as they appear in
# Ref09!D33:I33 and Ref11i!D:I (alphabetical, which SAVI's LOOKUPs rely on).
TM_ARRANGEMENTS = [
    "Contraflow",
    "Hardshoulder Closure",
    "Railway Possession",
    "Retaining Wall TM",
    "Shuttle Working / Lane Closure",
    "Waterway Possession",
]
TM_COLS = ["tm_contraflow", "tm_hardshoulder", "tm_railway", "tm_retaining_wall", "tm_shuttle_lane", "tm_waterway"]

CONDITION_CODES = ["1A", "2B", "2C", "2D", "2E", "3B", "3C", "3D", "3E",
                   "4B", "4C", "4D", "4E", "5B", "5C", "5D", "5E"]

# Structures sheet input columns A..AG (row 5 onwards); AH..AO are SAVI formulas
# derived from these and are recomputed by the engine.
STRUCTURE_INPUT_COLUMNS = [
    "structure_seq", "structure_id", "name", "reference", "structure_type", "number_in_group",
    "length_m", "width_m", "headroom_m", "spans", "location", "route_supported", "route_salted",
    "traffic", "obstacle_crossed", "obstacle_salted",
    "pf_heritage", "pf_2", "pf_3", "pf_4", "pf_5", "pf_6",
    "adj_heritage", "adj_conservation_area", "adj_environmentally_sensitive", "adj_river_coastal_wall",
    "adj_substandard", "adj_spare_1", "adj_spare_2", "adj_spare_3", "adj_spare_4", "adj_spare_5", "adj_spare_6",
]
ELEMENT_INPUT_COLUMNS = [
    "element_seq", "structure_id", "element_type", "material", "condition_code", "spray_proximity",
    "pf_heritage", "pf_2", "pf_3", "pf_4", "pf_5", "pf_6", "intervention_threshold",
]
PLANNED_INPUT_COLUMNS = ["element_seq", "structure_id", "structure_name", "element_type", "material",
                         "year_of_work", "maintenance_action", "total_cost", "condition_after"]
ROUTINE_INPUT_COLUMNS = ["year", "routine_maintenance", "general_inspections", "principal_inspections",
                         "structural_assessments", "other_regular_maintenance"]

# Adjustment factors in Structures!W..AO. The first 11 are user flags (inputs);
# the last 8 are derived by SAVI from other inventory fields.
ADJUSTMENT_FACTOR_COLUMNS = [
    "adj_heritage", "adj_conservation_area", "adj_environmentally_sensitive", "adj_river_coastal_wall",
    "adj_substandard", "adj_spare_1", "adj_spare_2", "adj_spare_3", "adj_spare_4", "adj_spare_5",
    "adj_spare_6", "adj_tunnel_over_400m", "adj_rural", "adj_route_unclassified", "adj_obstacle_footway",
    "adj_obstacle_watercourse_non_nav", "adj_obstacle_land_disused", "adj_obstacle_railway",
    "adj_obstacle_tenanted_business",
]


@dataclass
class SaviWorkbook:
    path: Path
    wb_values: openpyxl.Workbook

    @classmethod
    def open(cls, path: str | Path) -> "SaviWorkbook":
        path = Path(path)
        return cls(path, openpyxl.load_workbook(path, read_only=True, data_only=True))

    # -- low-level helpers -------------------------------------------------------
    def grid(self, sheet: str, rng: str) -> list[list]:
        """Return the cached values of an A1-style range as a list of rows."""
        a, b = rng.split(":")
        c0, r0 = _split(a)
        c1, r1 = _split(b)
        ws = self.wb_values[sheet]
        rows = ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True)
        return [list(r) for r in rows]

    def cell(self, sheet: str, ref: str):
        return self.grid(sheet, f"{ref}:{ref}")[0][0]

    def frame(self, sheet: str, rng: str, columns: list[str], drop_blank_on: str | None = None) -> pd.DataFrame:
        df = pd.DataFrame(self.grid(sheet, rng), columns=columns)
        if drop_blank_on:
            df = df[df[drop_blank_on].notna() & (df[drop_blank_on].astype(str).str.strip() != "")]
        return df.reset_index(drop=True)

    def last_row(self, sheet: str, column: str, start: int) -> int:
        """Last row index with a value in `column`, scanning down from `start`."""
        ws = self.wb_values[sheet]
        last = start - 1
        for i, (v,) in enumerate(ws.iter_rows(min_row=start, min_col=col(column), max_col=col(column),
                                              values_only=True), start):
            if v not in (None, ""):
                last = i
        return last


def _split(a1: str) -> tuple[int, int]:
    letters = "".join(ch for ch in a1 if ch.isalpha())
    digits = "".join(ch for ch in a1 if ch.isdigit())
    return col(letters), int(digits)


# ---------------------------------------------------------------------------
# Rule set
# ---------------------------------------------------------------------------

def read_ruleset(wb: SaviWorkbook) -> dict[str, pd.DataFrame]:
    """Extract SAVI's reference tables as tidy DataFrames keyed by table name."""
    t: dict[str, pd.DataFrame] = {}

    t["condition_codes"] = wb.frame("Ref00", "B13:H29",
                                    ["code", "ecs", "ext_s2", "ext_s3", "ext_s4", "ext_s5", "teod"]).fillna(0.0)
    t["locations"] = wb.frame("Ref00", "B3:C5", ["location", "exposure_factor"])
    t["spray_proximity"] = wb.frame("Ref00", "B8:C9", ["spray_proximity", "exposure_factor"])
    t["importance"] = wb.frame("Ref00", "B32:C35", ["importance", "eif"])

    # --- element types: Ref14 + Ref07&08 + Ref09 + Ref02i share one row order --------
    ref14 = wb.frame("Ref14", "A3:E76", ["element_type", "life", "importance", "eif", "excluded_from_valuation"])
    ref0708 = wb.frame("Ref07&08", "A3:K76", ["element_type", "size_formula", "unit", "size_note",
                                              "f_number", "f_length", "f_width", "f_headroom",
                                              "size_multiplier", "f_spans", "spans_offset"])
    ref09 = wb.frame("Ref09", "B34:I107", ["element_type", "tm_location"] + TM_COLS)
    exp_cols = [f"exp_flag_{i:02d}" for i in range(1, 22)]
    ref02i = wb.frame("Ref02i", "A3:W76", ["element_type", "exposure_scenario_id"] + exp_cols)
    for other in (ref0708, ref09, ref02i):
        if list(other["element_type"]) != list(ref14["element_type"]):
            raise ValueError("Element order differs between Ref14 and Ref07&08/Ref09/Ref02i")
    et = pd.concat([ref14,
                    ref0708.drop(columns="element_type"),
                    ref09.drop(columns="element_type"),
                    ref02i.drop(columns="element_type")], axis=1)
    et["excluded_from_valuation"] = et["excluded_from_valuation"].fillna(0).astype(int).astype(bool)
    for c in ["f_number", "f_length", "f_width", "f_headroom", "size_multiplier", "f_spans", "spans_offset",
              *TM_COLS, *exp_cols]:
        et[c] = pd.to_numeric(et[c], errors="coerce").fillna(0.0)
    et["tm_location"] = et["tm_location"].fillna("")
    et.insert(0, "element_index", range(1, len(et) + 1))
    t["element_types"] = et

    # Exposure encoding: weights per parameter (Ref02i!B77:W77) and the parameter names (row 2).
    hdr = wb.grid("Ref02i", "B2:W2")[0]
    weights = wb.grid("Ref02i", "B77:W77")[0]
    t["exposure_weights"] = pd.DataFrame({"position": range(len(weights)), "parameter": hdr,
                                          "weight": [int(w) for w in weights]})
    n = wb.last_row("Ref02ii", "A", 4)
    t["exposure_rules"] = wb.frame("Ref02ii", f"A4:B{n}", ["exposure_id", "exposure"])
    t["exposure_rules"]["exposure_id"] = t["exposure_rules"]["exposure_id"].astype("int64")

    # --- materials --------------------------------------------------------------------
    mats = wb.frame("Ref03&04i", "B3:B105", ["material"], drop_blank_on="material")
    by = wb.frame("Ref03&04i", "BY3:BY105", ["protected_by_finishes"])
    mats["protected_by_finishes"] = pd.to_numeric(by["protected_by_finishes"].iloc[:len(mats)],
                                                  errors="coerce").fillna(0).astype(int).values
    t["materials"] = mats

    elem_hdr = wb.grid("Ref03&04i", "C2:BX2")[0]
    lists = wb.grid("Ref03&04i", "C107:BX131")
    allowed = [(e, row[j]) for j, e in enumerate(elem_hdr) for row in lists
               if e and row[j] not in (None, "", 0)]
    t["element_materials"] = pd.DataFrame(allowed, columns=["element_type", "material"]).drop_duplicates()

    st_hdr = wb.grid("Ref11ii", "C2:W2")[0]
    st_lists = wb.grid("Ref11ii", "C3:W38")
    allowed = [(s, row[j]) for j, s in enumerate(st_hdr) for row in st_lists
               if s and row[j] not in (None, "", 0)]
    t["structure_elements"] = pd.DataFrame(allowed, columns=["structure_type", "element_type"]).drop_duplicates()

    det = wb.frame("Ref03&04ii", "A4:AA427",
                   ["material", "exposure"] + [f"frac_{c}" for c in CONDITION_CODES]
                   + ["mttf_with_global_factor", "notes", "material_factor", "mttf", "chk_material",
                      "chk_rule", "code", "mttf_base"], drop_blank_on="material")
    t["deterioration"] = det.drop(columns=["notes", "chk_material", "chk_rule", "code"])

    # --- maintenance actions and intervention rules ------------------------------------
    t["maintenance_actions"] = wb.frame(
        "Ref05&07", "A4:K96",
        ["action_id", "action", "unit", "cost_type", "fixed_rate", "constant_rate",
         "var_rate_s2", "var_rate_s3", "var_rate_s4", "var_rate_s5", "work_rate"], drop_blank_on="action")
    t["maintenance_actions_base"] = wb.frame(
        "Ref05&07", "N4:W96",
        ["action", "unit", "cost_type", "fixed_rate", "constant_rate",
         "var_rate_s2", "var_rate_s3", "var_rate_s4", "var_rate_s5", "work_rate"], drop_blank_on="action")
    n = wb.last_row("Ref06", "A", 4)
    t["intervention_rules"] = wb.frame(
        "Ref06", f"A4:I{n}",
        ["rule_id", "trigger_early", "trigger_mid", "trigger_late", "trigger_user",
         "material", "exposure", "action", "condition_after"])

    # --- costs add-ons and traffic management -------------------------------------------
    bands = []
    for kind, rng in (("fP", "B13:C16"), ("fO", "B20:C23"), ("fD", "B27:C30")):
        for thr, f in wb.grid("Ref09", rng):
            bands.append((kind, float(thr), float(f)))
    t["addon_bands"] = pd.DataFrame(bands, columns=["kind", "threshold", "factor"])
    names = wb.grid("Ref09", "D33:I33")[0]
    rates = wb.grid("Ref09", "D108:I108")[0]
    ids = wb.grid("Ref09", "D109:I109")[0]
    if names != TM_ARRANGEMENTS:
        raise ValueError(f"Unexpected TM arrangement order: {names}")
    t["tm_arrangements"] = pd.DataFrame({"arrangement_id": [int(i) for i in ids], "arrangement": names,
                                         "rate_per_hour": rates, "column": TM_COLS})

    # --- network factors -----------------------------------------------------------------
    traffic = wb.frame("Ref11i", "B10:D12", ["traffic", "tf", "exposure_factor"])
    delay = wb.frame("Ref10", "B25:C27", ["traffic", "daily_delay_rate"])
    t["traffic"] = traffic.merge(delay, on="traffic", how="left")
    t["obstacles"] = wb.frame("Ref11i", "B16:I32", ["obstacle", "obsf"] + TM_COLS)
    t["routes"] = wb.frame("Ref11i", "B36:I45", ["route", "rf"] + TM_COLS)
    t["priority_weights"] = pd.DataFrame(
        {"name": ["f1", "f2", "f3"], "value": [r[0] for r in wb.grid("Ref11i", "C4:C6")]})

    st = wb.frame("Ref11i", "B49:J69", ["structure_type", "stf", "avf_type", "avf_f_number", "avf_f_length",
                                         "avf_f_width", "avf_f_headroom", "tm_required", "dimension_multiplier"])
    grc = wb.frame("Ref12", "B5:K25", ["structure_type", "grc_unit_rate", "grc_note", "grc_f_number",
                                        "grc_f_length", "grc_f_width", "grc_f_headroom", "inflation_new_build",
                                        "structure_type_base", "grc_unit_rate_base"])
    if [s.lower() for s in st["structure_type"]] != [s.lower() for s in grc["structure_type"]]:
        raise ValueError("Structure type order differs between Ref11i and Ref12")
    st = pd.concat([st, grc.drop(columns=["structure_type", "inflation_new_build", "structure_type_base"])], axis=1)
    for c in ["avf_f_number", "avf_f_length", "avf_f_width", "avf_f_headroom", "tm_required",
              "grc_f_number", "grc_f_length", "grc_f_width", "grc_f_headroom"]:
        st[c] = pd.to_numeric(st[c], errors="coerce").fillna(0.0)
    t["structure_types"] = st
    t["avf_factors"] = wb.frame("Ref11i", "B73:E79", ["avf_type", "Local Authority",
                                                       "Overseeing (Trunk Road) Authority", "Transport for London"])

    t["traffic_restriction_codes"] = wb.frame("Ref10", "B5:D11", ["avf_type", "code_ecs4", "code_ecs5"])
    days = wb.grid("Ref10", "B18:H20")
    t["traffic_delay_days"] = pd.DataFrame(
        [(float(r[0]), code, float(r[code - 1])) for r in days for code in range(2, 8)],
        columns=["length_over_m", "restriction_code", "days"])

    # --- GRC adjustment factors (Structures!W3:AO3 values, labels in row 4) --------------
    vals = wb.grid("Structures", "W3:AO3")[0]
    labels = [str(x).strip() for x in wb.grid("Structures", "W4:AO4")[0]]
    t["adjustment_factors"] = pd.DataFrame({"column": ADJUSTMENT_FACTOR_COLUMNS, "label": labels,
                                            "factor": [float(v) for v in vals]})
    pf_vals = wb.grid("Structures", "Q3:V3")[0]
    epf_vals = wb.grid("Elements", "G3:L3")[0]
    t["priority_factor_weights"] = pd.DataFrame({
        "level": ["structure"] * 6 + ["element"] * 6,
        "column": ["pf_heritage", "pf_2", "pf_3", "pf_4", "pf_5", "pf_6"] * 2,
        "weight": [float(v) for v in pf_vals] + [float(v) for v in epf_vals]})

    # --- indexation, thresholds and constants ----------------------------------------------
    t["indexation"] = wb.frame("Ref15", "B4:F18", ["price_year", "maintenance_multiplier", "grc_multiplier",
                                                   "index_name", "tool_version"])
    t["constants"] = pd.DataFrame([
        ("sgr_sci_av_threshold", float(wb.cell("Ref15", "D33"))),
        ("sgr_sci_crit_threshold", float(wb.cell("Ref15", "D34"))),
        ("replace_vs_repair_grc_multiplier", float(wb.cell("Ref15", "D28"))),
        ("maintenance_multiplier", float(wb.cell("Ref05&07", "M3"))),
        ("grc_multiplier", float(wb.cell("Ref12", "I4"))),
        ("traffic_delay_extra_multiplier", float(wb.cell("Ref10", "D25"))),
        ("global_deterioration_factor", float(wb.cell("Ref03&04ii", "T2") or 0.0)),
    ], columns=["name", "value"])
    t["discount_rates"] = wb.frame("Ref15", "B39:D41", ["from_year", "to_year", "rate"])
    # SCI condition bands are hard-coded in SAVI's Y0 formulas (>=90, >=80, >=65, >=40, <40).
    t["condition_bands"] = pd.DataFrame({"band": ["Very Good", "Good", "Fair", "Poor", "Very Poor"],
                                         "sci_min": [90.0, 80.0, 65.0, 40.0, float("-inf")]})
    return t


# ---------------------------------------------------------------------------
# Inventory
# ---------------------------------------------------------------------------

@dataclass
class RawInventory:
    structures: pd.DataFrame
    elements: pd.DataFrame
    planned: pd.DataFrame
    routine: pd.DataFrame
    settings: dict


def read_inventory(wb: SaviWorkbook) -> RawInventory:
    """Read the user input sheets of a SAVI workbook."""
    n = wb.last_row("Structures", "B", 5)
    s = wb.frame("Structures", f"A5:AG{max(n, 5)}", STRUCTURE_INPUT_COLUMNS, drop_blank_on="structure_id")
    n = wb.last_row("Elements", "B", 5)
    e = wb.frame("Elements", f"A5:M{max(n, 5)}", ELEMENT_INPUT_COLUMNS, drop_blank_on="structure_id")
    e["intervention_threshold"] = e["intervention_threshold"].where(e["intervention_threshold"].notna(), None)
    n = wb.last_row("Planned", "A", 5)
    p = wb.frame("Planned", f"A5:I{max(n, 5)}", PLANNED_INPUT_COLUMNS, drop_blank_on="element_seq")
    r = wb.frame("Routine", "A5:F124", ROUTINE_INPUT_COLUMNS, drop_blank_on="year")
    settings = {
        "organisation": wb.cell("Main Page", "N27"),
        "authority_type": wb.cell("Main Page", "N29"),
        "price_year": wb.cell("Main Page", "AO46"),
        "intervention_strategy": wb.cell("Main Page", "BU31"),
        "input_date": str(wb.cell("Main Page", "AO18")),
        "years_long_term": wb.cell("Main Page", "CB16"),
        "budget_5yr_annual": wb.cell("Main Page", "BH19"),
        "budget_long_term": [list(r) for r in wb.grid("Main Page", "BW19:CA21")],
    }
    return RawInventory(s, e, p, r, settings)
